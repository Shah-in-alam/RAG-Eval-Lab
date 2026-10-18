import io
import zipfile

import pytest

from src.ingest import (MAX_BYTES, EmptyDocument, UnsafeUrl, UnsupportedFile, UploadStore, check_url_safe,
                        extract_text, normalize_share_url)


def make_docx() -> bytes:
    import docx
    d = docx.Document()
    d.add_heading("Quarterly report", level=1)
    d.add_paragraph("Revenue grew by 12 percent in the third quarter.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Region", "EMEA"
    buf = io.BytesIO(); d.save(buf); return buf.getvalue()


def make_pptx() -> bytes:
    from pptx import Presentation
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Roadmap"
    s.placeholders[1].text = "Launch the hybrid retriever in May."
    buf = io.BytesIO(); p.save(buf); return buf.getvalue()


def make_xlsx() -> bytes:
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Prices"
    ws.append(["item", "price"]); ws.append(["widget", 4.5])
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def make_pdf(text: str = "Hello PDF world about vector search") -> bytes:
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for i, o in enumerate(objs, 1):
        offsets.append(out.tell()); out.write(f"{i} 0 obj\n".encode() + o + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer << /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return out.getvalue()


def make_epub() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("mimetype", "application/epub+zip")
        z.writestr("OEBPS/ch1.xhtml", "<html><body><h1>Chapter One</h1><p>The whale was large.</p></body></html>")
    return buf.getvalue()


@pytest.mark.parametrize("name,data,expect", [
    ("notes.txt", "plain text body".encode(), "plain text body"),
    ("latin1.txt", "café menu".encode("cp1252"), "café menu"),
    ("readme.md", b"# Title\n\nSome *markdown* text", "== Title =="),
    ("page.html", b"<html><head><script>x=1</script></head><body><h2>Intro</h2><p>Hi there</p></body></html>", "== Intro =="),
    ("data.csv", b"name,age\nAda,36\n", "Ada"),
    ("data.json", b'{"name": "Ada", "skills": ["math"]}', "math"),
    ("doc.rtf", rb"{\rtf1\ansi Hello {\b bold} RTF}", "bold"),
])
def test_text_formats(name, data, expect):
    text = extract_text(name, data)
    assert expect in text
    assert "<script>" not in text and "x=1" not in text


def test_docx_keeps_headings_and_tables():
    text = extract_text("report.docx", make_docx())
    assert "== Quarterly report ==" in text and "Revenue grew" in text and "EMEA" in text


def test_pptx_slides_become_sections():
    text = extract_text("deck.pptx", make_pptx())
    assert "== Slide 1: Roadmap ==" in text and "hybrid retriever" in text


def test_xlsx_sheets_and_rows():
    text = extract_text("prices.xlsx", make_xlsx())
    assert "== Sheet: Prices ==" in text and "widget" in text and "4.5" in text


def test_pdf_text():
    text = extract_text("paper.pdf", make_pdf())
    assert "vector search" in text and "== Page 1 ==" in text


def test_epub_chapters():
    assert "The whale was large." in extract_text("book.epub", make_epub())


def test_unsupported_and_empty():
    with pytest.raises(UnsupportedFile):
        extract_text("virus.exe", b"MZ")
    with pytest.raises(EmptyDocument):
        extract_text("blank.txt", b"   \n ")
    with pytest.raises(EmptyDocument):
        extract_text("scan.pdf", make_pdf(""))


def test_share_links_become_direct_downloads():
    assert normalize_share_url("https://drive.google.com/file/d/ABC123/view?usp=sharing") == \
        "https://drive.google.com/uc?export=download&id=ABC123"
    assert normalize_share_url("https://docs.google.com/document/d/XYZ/edit") == \
        "https://docs.google.com/document/d/XYZ/export?format=docx"
    assert normalize_share_url("https://docs.google.com/spreadsheets/d/S1/edit#gid=0") == \
        "https://docs.google.com/spreadsheets/d/S1/export?format=xlsx"
    assert normalize_share_url("https://www.dropbox.com/s/k/file.pdf?dl=0") == "https://www.dropbox.com/s/k/file.pdf?dl=1"
    assert normalize_share_url("https://1drv.ms/b/s!AbCd").startswith("https://api.onedrive.com/v1.0/shares/u!")
    assert normalize_share_url("https://example.com/a.pdf") == "https://example.com/a.pdf"


def test_unsafe_urls_rejected():
    for bad in ["file:///etc/passwd", "ftp://example.com/x", "http://127.0.0.1:8000/api", "http://localhost/x",
                "http://10.0.0.5/doc.pdf", "http://169.254.169.254/latest/meta-data", "http://[::1]/x"]:
        with pytest.raises(UnsafeUrl):
            check_url_safe(bad)


def test_store_add_list_delete(tmp_path):
    store = UploadStore(tmp_path)
    v0 = store.version
    meta = store.add("My Report (final).docx", make_docx(), source="upload")
    assert meta["doc_id"].startswith("up_") and meta["kind"] == "docx" and meta["n_words"] > 5
    assert store.version != v0
    assert [m["doc_id"] for m in store.list()] == [meta["doc_id"]]
    assert "Revenue grew" in store.texts()[meta["doc_id"]]
    again = store.add("My Report (final).docx", make_docx(), source="upload")
    assert again["doc_id"] != meta["doc_id"]  # same name twice gets a distinct id
    assert store.delete(meta["doc_id"]) is True
    assert store.delete("up_nope") is False
    assert len(store.list()) == 1


def test_store_rejects_oversized(tmp_path):
    with pytest.raises(ValueError):
        UploadStore(tmp_path).add("big.txt", b"a" * (MAX_BYTES + 1), source="upload")
