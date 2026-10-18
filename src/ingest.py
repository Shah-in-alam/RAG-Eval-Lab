"""Turn user documents (files or links) into plain text the pipeline can index.

Supported: PDF, Word (.docx), OpenDocument text (.odt), PowerPoint (.pptx), Excel (.xlsx),
CSV, TXT, Markdown, HTML, RTF, EPUB and JSON. Section structure is kept as
`== Heading ==` lines so the heading chunker works on uploads too.

Links: plain file URLs plus Google Drive / Docs / Sheets / Slides, Dropbox and OneDrive
share links (converted to their direct-download form). Links must be public or
"anyone with the link"; private cloud files need the file to be downloaded first.
"""
from __future__ import annotations

import base64
import csv
import hashlib
import io
import ipaddress
import json
import re
import socket
import time
import uuid
import zipfile
from pathlib import Path
from urllib.parse import parse_qs, unquote, urlencode, urlparse, urlunparse

ROOT = Path(__file__).resolve().parent.parent
UPLOAD_DIR = ROOT / "data" / "uploads"
MAX_BYTES = 25 * 1024 * 1024  # per file
SUPPORTED = {".pdf", ".docx", ".odt", ".pptx", ".xlsx", ".csv", ".tsv", ".txt", ".md", ".markdown",
             ".html", ".htm", ".rtf", ".epub", ".json"}
LEGACY_HINT = {".doc": "Word 97-2003 (.doc)", ".ppt": "PowerPoint 97-2003 (.ppt)", ".xls": "Excel 97-2003 (.xls)"}
CONTENT_TYPE_EXT = {
    "application/pdf": ".pdf", "text/plain": ".txt", "text/markdown": ".md", "text/html": ".html", "text/csv": ".csv",
    "application/json": ".json", "application/rtf": ".rtf", "text/rtf": ".rtf", "application/epub+zip": ".epub",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": ".pptx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
    "application/vnd.oasis.opendocument.text": ".odt",
}


class UnsupportedFile(ValueError):
    pass


class EmptyDocument(ValueError):
    pass


class UnsafeUrl(ValueError):
    pass


# ---------------------------------------------------------------- extraction

def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _heading(text: str) -> str:
    text = " ".join(text.split())
    return f"\n== {text} ==\n" if text else ""


def _html_to_text(markup: str) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(markup, "html.parser")
    for tag in soup(["script", "style", "noscript", "nav", "footer", "header", "svg", "form"]):
        tag.decompose()
    for h in soup.find_all(["h1", "h2", "h3", "h4"]):
        h.replace_with(_heading(h.get_text(" ")))
    for br in soup.find_all(["p", "li", "tr", "div", "br"]):
        br.insert_after("\n")
    return soup.get_text()


def _markdown(text: str) -> str:
    return re.sub(r"(?m)^(#{1,6})\s+(.+?)\s*#*\s*$", lambda m: f"== {m.group(2)} ==", text)


def _docx(data: bytes) -> str:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(io.BytesIO(data))
    out = []
    for block in d.iter_inner_content():
        if isinstance(block, Paragraph):
            style = (block.style.name if block.style is not None else "") or ""
            if style.startswith("Heading") or style == "Title":
                out.append(_heading(block.text))
            elif block.text.strip():
                out.append(block.text)
        elif isinstance(block, Table):
            for row in block.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    out.append(" | ".join(dict.fromkeys(cells)))
    return "\n".join(out)


def _odt(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        xml = z.read("content.xml").decode("utf-8")
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(xml, "html.parser")
    out = []
    for el in soup.find_all(["text:h", "text:p"]):
        out.append(_heading(el.get_text(" ")) if el.name == "text:h" else el.get_text(" "))
    return "\n".join(out)


def _pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        out.append(_heading(f"Slide {i}: {title}" if title else f"Slide {i}"))
        for shape in slide.shapes:
            if shape is title_shape:
                continue
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    out.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append("Notes: " + slide.notes_slide.notes_text_frame.text)
    return "\n".join(out)


def _xlsx(data: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(_heading(f"Sheet: {ws.title}"))
        for row in ws.iter_rows(values_only=True):
            cells = [str(v) for v in row if v is not None and str(v).strip()]
            if cells:
                out.append(" | ".join(cells))
    wb.close()
    return "\n".join(out)


def _csv(text: str, delimiter: str) -> str:
    return "\n".join(" | ".join(c.strip() for c in row if c.strip()) for row in csv.reader(io.StringIO(text), delimiter=delimiter))


def _json(text: str) -> str:
    def walk(v, path=""):
        if isinstance(v, dict):
            for k, x in v.items():
                yield from walk(x, f"{path}.{k}" if path else str(k))
        elif isinstance(v, list):
            for i, x in enumerate(v):
                yield from walk(x, f"{path}[{i}]")
        else:
            yield f"{path}: {v}"
    return "\n".join(walk(json.loads(text)))


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        t = (page.extract_text() or "").strip()
        if t:
            pages.append(_heading(f"Page {i}") + t)
    if not pages:
        raise EmptyDocument("No selectable text in this PDF (it may be a scanned image; OCR is not supported).")
    return "\n".join(pages)


def _epub(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = z.namelist()
        order = [n for n in names if n.lower().endswith((".xhtml", ".html", ".htm"))]
        opf = next((n for n in names if n.lower().endswith(".opf")), None)
        if opf:  # follow the reading order (spine) when the package file exists
            from bs4 import BeautifulSoup

            pkg = BeautifulSoup(z.read(opf).decode("utf-8", "ignore"), "html.parser")
            base = opf.rsplit("/", 1)[0] + "/" if "/" in opf else ""
            hrefs = {i.get("id"): base + unquote(i.get("href", "")) for i in pkg.find_all("item")}
            spine = [hrefs.get(r.get("idref")) for r in pkg.find_all("itemref")]
            order = [n for n in spine if n in names] or order
        return "\n".join(_html_to_text(z.read(n).decode("utf-8", "ignore")) for n in order)


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext in LEGACY_HINT:
        raise UnsupportedFile(f"{LEGACY_HINT[ext]} files are not supported. Save the file in the newer format "
                              f"({ext}x) and upload that instead.")
    if ext not in SUPPORTED:
        raise UnsupportedFile(f"Unsupported file type {ext or '(none)'}. Supported: {', '.join(sorted(SUPPORTED))}")
    try:
        if ext == ".pdf":
            text = _pdf(data)
        elif ext == ".docx":
            text = _docx(data)
        elif ext == ".odt":
            text = _odt(data)
        elif ext == ".pptx":
            text = _pptx(data)
        elif ext == ".xlsx":
            text = _xlsx(data)
        elif ext == ".epub":
            text = _epub(data)
        else:
            raw = _decode(data)
            if ext in (".html", ".htm"):
                text = _html_to_text(raw)
            elif ext in (".md", ".markdown"):
                text = _markdown(raw)
            elif ext in (".csv", ".tsv"):
                text = _csv(raw, "\t" if ext == ".tsv" else ",")
            elif ext == ".json":
                text = _json(raw)
            elif ext == ".rtf":
                from striprtf.striprtf import rtf_to_text

                text = rtf_to_text(raw)
            else:
                text = raw
    except (EmptyDocument, UnsupportedFile):
        raise
    except Exception as e:  # corrupt or mislabelled files
        raise UnsupportedFile(f"Could not read {filename}: {type(e).__name__}: {e}") from e
    text = re.sub(r"[ \t\r\f\v]+\n", "\n", text.replace("\r\n", "\n"))
    text = re.sub(r"[ \t]{2,}", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not re.search(r"\w", text):
        raise EmptyDocument(f"{filename} contains no readable text.")
    return text + "\n"


# ---------------------------------------------------------------- links

def normalize_share_url(url: str) -> str:
    """Rewrite cloud share links to direct-download URLs."""
    p = urlparse(url.strip())
    host = p.netloc.lower()
    if host == "drive.google.com":
        m = re.search(r"/file/d/([^/]+)", p.path)
        file_id = m.group(1) if m else parse_qs(p.query).get("id", [None])[0]
        if file_id:
            return f"https://drive.google.com/uc?export=download&id={file_id}"
    if host == "docs.google.com":
        m = re.match(r"/(document|spreadsheets|presentation)/d/([^/]+)", p.path)
        if m:
            fmt = {"document": "docx", "spreadsheets": "xlsx", "presentation": "pptx"}[m.group(1)]
            return f"https://docs.google.com/{m.group(1)}/d/{m.group(2)}/export?format={fmt}"
    if host.endswith("dropbox.com"):
        q = parse_qs(p.query)
        q["dl"] = ["1"]
        return urlunparse(p._replace(query=urlencode({k: v[0] for k, v in q.items()})))
    if host in ("1drv.ms", "onedrive.live.com"):
        token = base64.urlsafe_b64encode(url.strip().encode()).decode().rstrip("=")
        return f"https://api.onedrive.com/v1.0/shares/u!{token}/root/content"
    if host.endswith("sharepoint.com") and "download=1" not in p.query:
        return urlunparse(p._replace(query=(p.query + "&" if p.query else "") + "download=1"))
    return url.strip()


def check_url_safe(url: str) -> str:
    """Allow only http(s) URLs whose host resolves to public addresses (no localhost / LAN / metadata)."""
    p = urlparse(url)
    if p.scheme not in ("http", "https") or not p.hostname:
        raise UnsafeUrl("Only http:// and https:// links are allowed.")
    try:
        infos = socket.getaddrinfo(p.hostname, p.port or (443 if p.scheme == "https" else 80))
    except socket.gaierror as e:
        raise UnsafeUrl(f"Could not resolve {p.hostname}.") from e
    for info in infos:
        ip = ipaddress.ip_address(info[4][0].split("%")[0])
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast or ip.is_unspecified:
            raise UnsafeUrl(f"Links to local or private network addresses are not allowed ({p.hostname}).")
    return url


def _filename_from(resp, url: str) -> str:
    cd = resp.headers.get("content-disposition", "")
    m = re.search(r"filename\*=(?:UTF-8'')?([^;]+)", cd, re.I) or re.search(r'filename="?([^";]+)"?', cd, re.I)
    name = unquote(m.group(1).strip()) if m else unquote(Path(urlparse(url).path).name)
    ctype = resp.headers.get("content-type", "").split(";")[0].strip().lower()
    if Path(name).suffix.lower() not in SUPPORTED and ctype in CONTENT_TYPE_EXT:
        name = (name or "download") + CONTENT_TYPE_EXT[ctype]
    return name or "download.txt"


def fetch_url(url: str, max_bytes: int = MAX_BYTES, timeout: float = 30.0) -> tuple[str, bytes]:
    """Download a (share) link safely: checks every redirect hop and caps the size."""
    import httpx

    current = normalize_share_url(url)
    with httpx.Client(timeout=timeout, follow_redirects=False,
                      headers={"User-Agent": "RAGEvalLab/0.1 (document import)"}) as client:
        for _ in range(6):
            check_url_safe(current)
            with client.stream("GET", current) as resp:
                if resp.status_code in (301, 302, 303, 307, 308) and "location" in resp.headers:
                    current = str(resp.url.join(resp.headers["location"]))
                    continue
                if resp.status_code >= 400:
                    raise ValueError(f"The link returned HTTP {resp.status_code}. Is it shared publicly?")
                size = int(resp.headers.get("content-length") or 0)
                if size > max_bytes:
                    raise ValueError(f"File is larger than {max_bytes // (1024 * 1024)} MB.")
                buf = bytearray()
                for chunk in resp.iter_bytes():
                    buf.extend(chunk)
                    if len(buf) > max_bytes:
                        raise ValueError(f"File is larger than {max_bytes // (1024 * 1024)} MB.")
                name = _filename_from(resp, current)
                if "drive.google.com" in current and name.endswith(".html") and b"confirm" in bytes(buf[:20000]):
                    raise ValueError("Google Drive asked for a virus-scan confirmation (large file). "
                                     "Download it and upload the file instead.")
                return name, bytes(buf)
    raise ValueError("Too many redirects.")


# ---------------------------------------------------------------- storage

def _slug(name: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "_", Path(name).stem.lower()).strip("_")
    return s[:40] or "document"


class UploadStore:
    """User documents on disk: <id>.txt (extracted text) + <id>.json (metadata)."""

    def __init__(self, root: str | Path = UPLOAD_DIR) -> None:
        self.root = Path(root)

    @property
    def version(self) -> str:
        ids = sorted(p.stem for p in self.root.glob("*.json")) if self.root.exists() else []
        return hashlib.sha1("|".join(ids).encode()).hexdigest()[:12]

    def add(self, filename: str, data: bytes, source: str, url: str | None = None) -> dict:
        if len(data) > MAX_BYTES:
            raise ValueError(f"{filename} is larger than {MAX_BYTES // (1024 * 1024)} MB.")
        text = extract_text(filename, data)
        doc_id = f"up_{_slug(filename)}_{uuid.uuid4().hex[:6]}"
        meta = {"doc_id": doc_id, "title": Path(filename).stem, "filename": Path(filename).name,
                "kind": Path(filename).suffix.lower().lstrip("."), "size_bytes": len(data),
                "n_words": len(text.split()), "n_chars": len(text), "source": source, "url": url,
                "uploaded_at": time.strftime("%Y-%m-%dT%H:%M:%S")}
        self.root.mkdir(parents=True, exist_ok=True)
        (self.root / f"{doc_id}.txt").write_text(text, encoding="utf-8")
        (self.root / f"{doc_id}.json").write_text(json.dumps(meta, indent=1), encoding="utf-8")
        return meta

    def list(self) -> list[dict]:
        if not self.root.exists():
            return []
        metas = [json.loads(p.read_text(encoding="utf-8")) for p in self.root.glob("*.json")]
        return sorted(metas, key=lambda m: m["uploaded_at"], reverse=True)

    def texts(self) -> dict[str, str]:
        return {m["doc_id"]: (self.root / f"{m['doc_id']}.txt").read_text(encoding="utf-8") for m in self.list()}

    def get(self, doc_id: str) -> dict | None:
        p = self.root / f"{doc_id}.json"
        return json.loads(p.read_text(encoding="utf-8")) if re.fullmatch(r"up_[a-z0-9_]+", doc_id) and p.exists() else None

    def delete(self, doc_id: str) -> bool:
        if not self.get(doc_id):
            return False
        for ext in (".txt", ".json"):
            (self.root / f"{doc_id}{ext}").unlink(missing_ok=True)
        return True
